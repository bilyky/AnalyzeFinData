"""
Tests for the AI-buildout supply-chain watchlist (aether/ai_buildout.py).

Pure functions over fixture SEC / OHLCV JSON; no network. Covers theme bucketing,
quarterly revenue YoY (tag choice, restatements, look-ahead), RPO YoY, 8-K item
1.01 counting, placeholder-bar skipping, CMF, and the watch-score ranking.
"""
import json
import os
import sys
import tempfile
import unittest
import warnings
from pathlib import Path
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import requests
import urllib3

from aether import ai_buildout as ab


def _q(start, end, val, filed=None):
    return {"start": start, "end": end, "val": val, "filed": filed or end, "form": "10-Q"}


def _facts(tags):
    return {"facts": {"us-gaap": {t: {"units": {"USD": v}} for t, v in tags.items()}}}


def _research_book(path, symbols):
    """Minimal state_of_the_day.xlsx: a Research sheet with symbols in column D."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Research"
    ws.append(["A", "B", "C", "Symbol", "Industry"])
    for s in symbols:
        ws.append([None, None, None, s, "x"])
    wb.create_sheet("Other").append(["ignored", None, None, "NOPE"])
    wb.save(path)


def _bar(c, h=None, l=None, v=1000, **extra):
    h = c + 1 if h is None else h
    l = c - 1 if l is None else l
    return {"1. open": str(c), "2. high": str(h), "3. low": str(l),
            "4. close": str(c), "5. volume": str(v), **extra}


class TestBucket(unittest.TestCase):
    def test_seed_wins_over_sic(self):
        self.assertEqual(ab.bucket_for("sei", sic=1389), "power")
        self.assertEqual(ab.bucket_for("VRT", sic=4911), "cooling")

    def test_sic_mapping_and_unknowns(self):
        self.assertEqual(ab.bucket_for("XYZ", sic="4911"), "power")
        self.assertEqual(ab.bucket_for("XYZ", sic=1731), "construction")
        self.assertIsNone(ab.bucket_for("XYZ", sic=7372))
        self.assertIsNone(ab.bucket_for("TECK", sic=1400))   # mining: too broad
        self.assertIsNone(ab.bucket_for("RUN", sic=1700))    # general trades: too broad
        self.assertIsNone(ab.bucket_for("XYZ", sic=None))
        self.assertIsNone(ab.bucket_for("XYZ", sic="n/a"))


class TestRevenueYoy(unittest.TestCase):
    def test_quarterly_yoy_ignores_ytd_durations(self):
        facts = _facts({"Revenues": [
            _q("2025-04-01", "2025-06-30", 100),
            _q("2026-01-01", "2026-06-30", 999),   # 6-month YTD, must be ignored
            _q("2026-04-01", "2026-06-30", 150),
        ]})
        self.assertEqual(ab.revenue_yoy(facts), ("2026-06-30", 50.0))

    def test_picks_tag_with_latest_quarter(self):
        facts = _facts({
            "Revenues": [_q("2018-10-01", "2018-12-31", 5), _q("2019-10-01", "2019-12-31", 6)],
            "RevenueFromContractWithCustomerExcludingAssessedTax": [
                _q("2025-04-01", "2025-06-30", 200), _q("2026-04-01", "2026-06-30", 180)],
        })
        self.assertEqual(ab.revenue_yoy(facts), ("2026-06-30", -10.0))

    def test_restatement_uses_latest_filed(self):
        facts = _facts({"Revenues": [
            _q("2025-04-01", "2025-06-30", 100, filed="2025-08-01"),
            _q("2025-04-01", "2025-06-30", 120, filed="2026-08-01"),   # restated
            _q("2026-04-01", "2026-06-30", 150, filed="2026-08-01"),
        ]})
        self.assertEqual(ab.revenue_yoy(facts)[1], 25.0)

    def test_no_look_ahead(self):
        facts = _facts({"Revenues": [
            _q("2025-01-01", "2025-03-31", 100), _q("2026-01-01", "2026-03-31", 110),
            _q("2026-04-01", "2026-06-30", 500),
        ]})
        self.assertEqual(ab.revenue_yoy(facts, as_of="2026-05-15"), ("2026-03-31", 10.0))

    def test_point_in_time_uses_filed_date(self):
        facts = _facts({"Revenues": [
            _q("2025-04-01", "2025-06-30", 100, filed="2025-08-01"),
            _q("2025-04-01", "2025-06-30", 120, filed="2026-08-01"),   # restated later
            _q("2026-04-01", "2026-06-30", 150, filed="2026-08-05"),
        ]})
        # quarter ended 6/30 but not filed until 8/5: not known on 7/15
        self.assertEqual(ab.revenue_yoy(facts, as_of="2026-07-15"), ("2025-06-30", None))
        # 8/3: the restatement (filed 8/1) is known, the new quarter is not
        self.assertEqual(ab.revenue_yoy(facts, as_of="2026-08-03"), ("2025-06-30", None))
        # 8/10: new quarter vs the restated year-ago value
        self.assertEqual(ab.revenue_yoy(facts, as_of="2026-08-10"), ("2026-06-30", 25.0))
        # 2025-09-01: only the original filing of the 2025 quarter is known
        old = _facts({"Revenues": [_q("2024-04-01", "2024-06-30", 80, filed="2024-08-01"),
                                   _q("2025-04-01", "2025-06-30", 100, filed="2025-08-01"),
                                   _q("2025-04-01", "2025-06-30", 120, filed="2026-08-01")]})
        self.assertEqual(ab.revenue_yoy(old, as_of="2025-09-01"), ("2025-06-30", 25.0))

    def test_rpo_point_in_time(self):
        facts = _facts({"RevenueRemainingPerformanceObligation": [
            {"end": "2025-06-30", "val": 20e9, "filed": "2025-08-01"},
            {"end": "2026-06-30", "val": 30e9, "filed": "2026-08-01"},
        ]})
        self.assertEqual(ab.rpo_yoy(facts, as_of="2026-07-31"), ("2025-06-30", None))
        self.assertEqual(ab.rpo_yoy(facts, as_of="2026-08-01"), ("2026-06-30", 50.0))

    def test_missing_year_ago_gives_none(self):
        facts = _facts({"Revenues": [_q("2026-04-01", "2026-06-30", 150)]})
        self.assertEqual(ab.revenue_yoy(facts), ("2026-06-30", None))
        self.assertEqual(ab.revenue_yoy({}), (None, None))


class TestRpoYoy(unittest.TestCase):
    def test_instant_yoy_skips_zero(self):
        facts = _facts({"RevenueRemainingPerformanceObligation": [
            {"end": "2019-12-31", "val": 0, "filed": "2020-02-01"},
            {"end": "2025-06-30", "val": 20e9, "filed": "2025-08-01"},
            {"end": "2026-06-30", "val": 30e9, "filed": "2026-08-01"},
        ]})
        self.assertEqual(ab.rpo_yoy(facts), ("2026-06-30", 50.0))
        self.assertEqual(ab.rpo_yoy({}), (None, None))


class TestMaterialAgreements(unittest.TestCase):
    SUB = {"filings": {"recent": {
        "form": ["8-K", "8-K", "10-Q", "8-K", "8-K"],
        "filingDate": ["2026-10-01", "2026-09-22", "2026-08-05", "2026-09-08", "2026-05-01"],
        "items": ["1.01,2.03,9.01", "7.01,8.01", "", "1.01,2.01", "1.01"],
    }}}

    def test_counts_item_101_in_window(self):
        self.assertEqual(ab.material_agreements(self.SUB, "2026-10-02"),
                         (2, ["2026-10-01", "2026-09-08"]))

    def test_excludes_future_filings(self):
        self.assertEqual(ab.material_agreements(self.SUB, "2026-09-30")[0], 1)

    def test_item_match_is_exact(self):
        sub = {"filings": {"recent": {"form": ["8-K"], "filingDate": ["2026-09-01"],
                                      "items": ["11.01"]}}}
        self.assertEqual(ab.material_agreements(sub, "2026-10-01")[0], 0)


class TestPriceSignals(unittest.TestCase):
    def test_real_bars_skip_placeholders_and_future(self):
        ts = {
            "2026-09-24": _bar(100),
            "2026-09-25": _bar(101, v=0),                   # zero volume
            "2026-09-26": _bar(102, h=102, l=102),          # flat
            "2026-09-27": _bar(103, provisional=True),      # provisional
            "2026-09-28": _bar(104),
            "2026-09-29": _bar(105),                        # after as_of
        }
        bars = ab._real_bars(ts, "2026-09-28")
        self.assertEqual([b[0] for b in bars], ["2026-09-24", "2026-09-28"])

    def test_cmf_bounds(self):
        up = [("d", 11.0, 9.0, 11.0, 100.0)] * 20     # closes at the high
        down = [("d", 11.0, 9.0, 9.0, 100.0)] * 20    # closes at the low
        self.assertEqual(ab.chaikin_money_flow(up), 1.0)
        self.assertEqual(ab.chaikin_money_flow(down), -1.0)
        self.assertIsNone(ab.chaikin_money_flow(up[:5]))

    def test_return_pct(self):
        bars = [("d", 0, 0, 100.0, 1)] + [("d", 0, 0, 110.0, 1)] * 60
        self.assertAlmostEqual(ab.return_pct(bars, 60), 10.0)
        self.assertIsNone(ab.return_pct(bars[:10], 60))


class TestWatchScore(unittest.TestCase):
    def test_strong_and_weak(self):
        strong = {"revenue_yoy": 40, "rpo_yoy": 30, "agreements_90d": 3,
                  "rs_60d": 15, "cmf_20": 0.2}
        weak = {"revenue_yoy": -5, "rpo_yoy": -2, "agreements_90d": 0,
                "rs_60d": -8, "cmf_20": -0.2}
        self.assertEqual(ab.watch_score(strong), 9)
        self.assertEqual(ab.watch_score(weak), -4)

    def test_build_row_flags_suspect_rpo(self):
        facts = _facts({"RevenueRemainingPerformanceObligation": [
            {"end": "2025-06-30", "val": 7e6, "filed": "2025-08-01"},
            {"end": "2026-06-30", "val": 258e6, "filed": "2026-08-01"},
        ]})
        row = ab.build_row("AES", "power", {}, facts, None, [], "2026-10-02")
        self.assertTrue(row["rpo_suspect"])
        self.assertFalse(ab.build_row("X", "power", {}, {}, None, [], "2026-10-02")["rpo_suspect"])

    def test_ohlcv_reads_cache_dir_not_data_dir(self):
        with tempfile.TemporaryDirectory() as cache, tempfile.TemporaryDirectory() as data:
            (Path(cache) / "Symbol_full").mkdir()
            (Path(cache) / "Symbol_full" / "VRT_daily.json").write_text(
                json.dumps({"Time Series (Daily)": {"2026-01-02": _bar(10)}}), encoding="utf-8")
            with mock.patch.dict(os.environ, {"AETHER_CACHE_DIR": cache, "AETHER_DATA_DIR": data}):
                self.assertIn("2026-01-02", ab._load_ohlcv("VRT"))

    def test_suspect_rpo_jump_not_scored(self):
        base = {"revenue_yoy": None, "agreements_90d": 0, "rs_60d": None, "cmf_20": None}
        self.assertEqual(ab.watch_score({**base, "rpo_yoy": 3585.7}), 0)
        self.assertEqual(ab.watch_score({**base, "rpo_yoy": 75.0}), 2)

    def test_stale_figures_not_scored(self):
        base = {"revenue_yoy": 362.6, "rpo_yoy": None, "agreements_90d": 0,
                "rs_60d": None, "cmf_20": None}
        self.assertEqual(ab.watch_score(base), 2)
        self.assertEqual(ab.watch_score({**base, "stale": ["revenue_yoy"]}), 0)

    def test_build_row_marks_old_quarters_stale(self):
        facts = _facts({"Revenues": [_q("2024-10-01", "2024-12-31", 100),
                                     _q("2025-10-01", "2025-12-31", 400)]})
        row = ab.build_row("ARBE", "eyes", {}, facts, None, [], "2026-10-02")
        self.assertEqual(row["stale"], ["revenue_yoy"])
        self.assertEqual(row["watch_score"], 0)
        fresh = ab.build_row("ARBE", "eyes", {}, facts, None, [], "2026-03-01")
        self.assertEqual(fresh["stale"], [])
        self.assertEqual(fresh["watch_score"], 2)

    def test_missing_data_is_neutral(self):
        self.assertEqual(ab.watch_score({"revenue_yoy": None, "rpo_yoy": None,
                                         "agreements_90d": 0, "rs_60d": None,
                                         "cmf_20": None}), 0)


class TestCache(unittest.TestCase):
    def test_read_cache_ignores_age_and_missing(self):
        with tempfile.TemporaryDirectory() as d, mock.patch.dict(os.environ, {"AETHER_DATA_DIR": d}):
            (Path(d) / "edgar_cache").mkdir()
            (Path(d) / "edgar_cache" / "sub_1.json").write_text('{"sic": 4911}', encoding="utf-8")
            os.utime(Path(d) / "edgar_cache" / "sub_1.json", (0, 0))   # very old
            self.assertEqual(ab._read_cache("sub_1.json"), {"sic": 4911})
            self.assertIsNone(ab._read_cache("sub_2.json"))


class _Resp:
    def __init__(self, status, data=None):
        self.status_code = status
        self._data = data

    def json(self):
        return self._data

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


class TestEdgarBlock(unittest.TestCase):
    def setUp(self):
        self._d = tempfile.TemporaryDirectory()
        self._env = mock.patch.dict(os.environ, {"AETHER_DATA_DIR": self._d.name})
        self._env.start()
        ab._consecutive_403[0] = 0
        ab._blocked[0] = False
        self._gap = mock.patch.object(ab, "_REQUEST_GAP_S", 0)
        self._gap.start()

    def tearDown(self):
        self._gap.stop()
        self._env.stop()
        self._d.cleanup()
        ab._consecutive_403[0] = 0
        ab._blocked[0] = False

    def test_stops_requesting_after_consecutive_403s(self):
        http = mock.Mock(return_value=_Resp(403))
        with mock.patch.object(ab, "_http_get", http):
            for i in range(10):
                self.assertIsNone(ab._get_json(f"u{i}", f"x{i}.json"))
        self.assertEqual(http.call_count, ab._BLOCK_AFTER_403S)
        self.assertTrue(ab._blocked[0])

    def test_success_resets_403_count(self):
        responses = [_Resp(403), _Resp(403), _Resp(200, {"ok": 1}), _Resp(403), _Resp(403)]
        with mock.patch.object(ab, "_http_get", side_effect=responses):
            for i in range(5):
                ab._get_json(f"u{i}", f"y{i}.json")
        self.assertFalse(ab._blocked[0])

    def test_failed_refresh_falls_back_to_stale_cache(self):
        cache = Path(self._d.name) / "edgar_cache"
        cache.mkdir()
        (cache / "sub_9.json").write_text('{"sic": 4911}', encoding="utf-8")
        os.utime(cache / "sub_9.json", (0, 0))   # older than any TTL
        with mock.patch.object(ab, "_http_get", return_value=_Resp(403)):
            self.assertEqual(ab._get_json("u", "sub_9.json"), {"sic": 4911})
        ab._blocked[0] = True
        with mock.patch.object(ab, "_http_get") as http:
            self.assertEqual(ab._get_json("u", "sub_9.json"), {"sic": 4911})
            http.assert_not_called()


class TestTlsFallback(unittest.TestCase):
    def tearDown(self):
        ab._tls_fallback[0] = False

    def test_insecure_warning_silenced_only_for_the_call(self):
        def fake_get(url, headers=None, timeout=None, verify=True):
            if verify:
                raise requests.exceptions.SSLError("intercepted")
            warnings.warn("unverified", urllib3.exceptions.InsecureRequestWarning, stacklevel=2)
            return "resp"

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            before = list(warnings.filters)
            with mock.patch.object(ab.requests, "get", side_effect=fake_get):
                self.assertEqual(ab._http_get("https://x"), "resp")
            # filters unchanged right after the call: nothing added process-wide
            self.assertEqual(warnings.filters, before)
            self.assertFalse([w for w in caught
                              if issubclass(w.category, urllib3.exceptions.InsecureRequestWarning)])


class TestUniverse(unittest.TestCase):
    def test_reads_research_sheet_column_d(self):
        with tempfile.TemporaryDirectory() as d, mock.patch.dict(os.environ, {"AETHER_DATA_DIR": d}):
            _research_book(Path(d) / "state_of_the_day.xlsx", ["vrt", "ETN", None, " ETN ", "PWR"])
            self.assertEqual(ab.load_universe(), ["VRT", "ETN", "PWR"])

    def test_missing_workbook_raises(self):
        with tempfile.TemporaryDirectory() as d, mock.patch.dict(os.environ, {"AETHER_DATA_DIR": d}):
            with self.assertRaises(FileNotFoundError):
                ab.load_universe()


class TestScan(unittest.TestCase):
    def test_scan_classifies_and_ranks(self):
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "Symbol_full").mkdir()
            _research_book(Path(d) / "state_of_the_day.xlsx",
                           ["VRT", "MSFT", "UTIL", "GONE", "SPY"])
            subs = {1: {"sic": 3585, "name": "Vertiv"}, 2: {"sic": 7372, "name": "Microsoft"},
                    3: {"sic": 4911, "name": "Some Utility"}}
            cik = {"VRT": 1, "MSFT": 2, "UTIL": 3, "GONE": 4}   # 4: fetch fails
            facts = {1: _facts({"Revenues": [_q("2025-04-01", "2025-06-30", 100),
                                             _q("2026-04-01", "2026-06-30", 140)]}),
                     3: {}}
            with mock.patch.dict(os.environ, {"AETHER_DATA_DIR": d, "AETHER_CACHE_DIR": d}), \
                 mock.patch.object(ab, "ticker_cik_map", return_value=cik), \
                 mock.patch.object(ab, "submissions", side_effect=subs.get), \
                 mock.patch.object(ab, "company_facts", side_effect=facts.get):
                failed = []
                rows = ab.scan("2026-10-02", universe=ab.load_universe(), failed=failed)
                out = ab.save(rows, "2026-10-02", failed=failed)
                saved = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual([r["symbol"] for r in rows], ["VRT", "UTIL"])   # MSFT not themed
        self.assertEqual(rows[0]["revenue_yoy"], 40.0)
        self.assertEqual(rows[0]["watch_score"], 2)
        self.assertEqual(saved["as_of"], "2026-10-02")
        self.assertEqual(saved["fetch_failed"], ["GONE"])

    def test_scan_raises_without_ticker_map(self):
        with mock.patch.object(ab, "ticker_cik_map", return_value={}):
            with self.assertRaises(RuntimeError):
                ab.scan("2026-10-02", universe=["VRT"])



class TestThemes(unittest.TestCase):
    def test_robot_vision_buckets_are_seed_only(self):
        self.assertEqual(ab.bucket_for("OUST", theme="robot_vision"), "eyes")
        self.assertEqual(ab.bucket_for("LSCC", theme="robot_vision"), "chips")
        self.assertEqual(ab.bucket_for("MBLY", theme="robot_vision"), "software")
        # no SIC fallback: a random chip maker is not pulled in
        self.assertIsNone(ab.bucket_for("XYZ", sic=3674, theme="robot_vision"))
        # themes do not leak into each other
        self.assertIsNone(ab.bucket_for("OUST", theme="ai_buildout"))
        self.assertIsNone(ab.bucket_for("VRT", theme="robot_vision"))

    def test_every_seed_bucket_is_declared(self):
        for name, t in ab.THEMES.items():
            self.assertTrue(set(t["seed"].values()) <= set(t["buckets"]), name)
            self.assertTrue(set(t["sic"].values()) <= set(t["buckets"]), name)

    def test_unknown_theme_raises(self):
        with self.assertRaises(ValueError):
            ab.bucket_for("OUST", theme="nope")
        with self.assertRaises(ValueError):
            ab.output_path("../evil")

    def test_seed_only_scan_skips_universe_and_saves_per_theme(self):
        seen = []

        def sub(cik):
            seen.append(cik)
            return {"sic": 3674, "name": f"c{cik}"}

        cik = {s: i for i, s in enumerate(ab.ROBOT_VISION_SEED, start=1)}
        cik["UNIV"] = 999
        with tempfile.TemporaryDirectory() as d:
            with mock.patch.dict(os.environ, {"AETHER_DATA_DIR": d}),                  mock.patch.object(ab, "ticker_cik_map", return_value=cik),                  mock.patch.object(ab, "submissions", side_effect=sub),                  mock.patch.object(ab, "company_facts", return_value={}):
                rows = ab.scan("2026-10-02", universe=["UNIV"], theme="robot_vision")
                out = ab.save(rows, "2026-10-02", theme="robot_vision")
                self.assertEqual(out.name, "robot_vision_watch.json")
                self.assertEqual(ab.load_latest(theme="robot_vision")["theme"], "robot_vision")
                self.assertIsNone(ab.load_latest(theme="ai_buildout"))
        self.assertNotIn(999, seen)   # universe symbol never fetched
        self.assertEqual({r["symbol"] for r in rows}, set(ab.ROBOT_VISION_SEED))

if __name__ == "__main__":
    unittest.main()
