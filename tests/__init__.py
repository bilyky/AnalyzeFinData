# Market-data caches (Data/Symbol, Data/Symbol_full) -> temp, FIRST, before any module
# computes its cache constants at import (aether.paths.cache_dir() reads AETHER_CACHE_DIR).
# A suite run then never reads or writes the real caches, whether the checkout's Data/
# holds them directly or through a junction. Unconditional, like the ledger guard.
# Guarded by tests/test_log_hermetic.py::TestMarketDataCachesRedirected.
import os as _os_early
import tempfile as _tempfile_early

_test_cache_dir = _tempfile_early.TemporaryDirectory(ignore_cleanup_errors=True)
_os_early.environ["AETHER_CACHE_DIR"] = _test_cache_dir.name

# Globally mock notify.send_email during all unit test executions
# to completely prevent test email spam and keep production code clean.
import unittest.mock as _mock
import aether.notify as _notify
import notify as _root_notify

_notify.send_email = _mock.MagicMock(return_value=True)
_root_notify.send_email = _mock.MagicMock(return_value=True)

# Globally redirect all test log outputs to a temporary directory
# to prevent tests from writing to or polluting production logs.
import tempfile
from pathlib import Path
import aether.paths as _paths
import aether.logger

_test_log_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
aether.logger._LOG_DIR = Path(_test_log_dir.name)
# aether.logger opens its file handlers at IMPORT time (module-level
# `log = get_logger("aether")`), so re-pointing _LOG_DIR alone is too late — the
# handlers already target the repo's Data/logs. Close them and rebuild the chain
# under the temp dir. Guarded by tests/test_log_hermetic.py.
import logging as _logging

_aether_root_logger = _logging.getLogger("aether")
for _h in list(_aether_root_logger.handlers):
    _aether_root_logger.removeHandler(_h)
    _h.close()
aether.logger._initialised = False
aether.logger._init()
# daily_task.py attaches its own FileHandler (repo-root daily_task.log) at import,
# but only when its logger has no handlers yet — pre-seed one under the temp dir.
_logging.getLogger("aether.daily_task").addHandler(
    _logging.FileHandler(Path(_test_log_dir.name) / "daily_task.log", encoding="utf-8", delay=True))

# ---------------------------------------------------------------------------
# Hermeticity guard — no test may touch prod state, contact a live host, or
# open a real browser. This whole block is one switch: AETHER_LIVE_TESTS=1
# disables it for the explicitly live-gated contract tests (test_live_api_contract,
# test_game_pricing, test_etrade_live), which INTEND to touch the real token files
# and reach the real broker over ban-safe HTTP. Everything below stays gated
# together so a live run sees real files AND a real network, and a hermetic run
# sees neither.
# ---------------------------------------------------------------------------
# A plain `python -m unittest discover tests` must NEVER reach E*TRADE, Chaikin,
# Google, or any external API, nor read/write the production auth-state files. On
# 2026-08-18 an unmocked requalify_symbol("AAPL") in the suite fell through to a
# live E*TRADE re-auth + a real Chaikin browser login, tripping Akamai Bot Manager
# and driving toward an IP ban.
import os as _os
import socket as _socket
import re as _re
import subprocess as _subprocess
import importlib as _importlib
import aether.etrade as _etrade
import aether.instruments as _instruments

# playwright may be absent in some environments; import it dynamically (no
# top-of-block import statement, so the pre-commit import linter stays happy)
# and fail soft to None if it isn't installed.
try:
    _pw_sync = _importlib.import_module("playwright.sync_api")
except Exception:
    _pw_sync = None

# ---------------------------------------------------------------------------
# Prod-workbook guard — no test may WRITE the production state_of_the_day.xlsx.
# ---------------------------------------------------------------------------
# The workbook is gitignored live-portfolio state. powergauge.check_from_xls and
# the pipeline savers call wb.save(XLSX_FILE); a save-path test would corrupt
# real trading state. Point every module's XLSX_FILE at a throwaway temp book:
# .exists() and load_workbook still work, every save lands in temp. Every test
# that reads the workbook mocks load_workbook with its own in-memory book, so the
# temp only has to exist — it deliberately does NOT mirror prod (no real data in
# temp, no per-import copy). Unconditional: the live tier opts into the real
# network, never into writing prod state. Per-test overrides (test_breadth_filter)
# still win; they default to temp now.
import openpyxl as _openpyxl

_test_xlsx_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
_test_xlsx_path = Path(_test_xlsx_dir.name) / "state_of_the_day.xlsx"
_seed_wb = _openpyxl.Workbook()
_seed_wb.active.title = "Research"
_seed_wb.save(_test_xlsx_path)

_test_game_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
_test_game_file = Path(_test_game_dir.name) / "ai_portfolio_game.json"

for _mod_name in ("ai_portfolio_game", "powergauge", "workbook_read", "autonomous_pipeline", "bootstrap_dna"):
    try:
        _mod = _importlib.import_module(_mod_name)
    except Exception:
        continue
    # autonomous_pipeline.log() also appends to its own legacy run log
    # (Data/autonomous_run.log, read by watchdog and the dashboard) — redirect it too.
    if hasattr(_mod, "LOG_FILE_PATH"):
        _mod.LOG_FILE_PATH = Path(_test_log_dir.name) / "autonomous_run.log"
    # Game state: ai_portfolio_game is the only writer (save_game, load_game's backup
    # restore). Its backups MUST move with it — save_game prunes GAME_BACKUP_DIR to the
    # last 15, so redirecting the file alone would delete real backups. Readers
    # (workbook_read / bootstrap_dna GAME_FILE) point at the same temp file.
    if hasattr(_mod, "AI_GAME_FILE"):
        _mod.AI_GAME_FILE = _test_game_file
        _mod.GAME_BACKUP_DIR = Path(_test_game_dir.name) / "Backup" / "Game"
    if hasattr(_mod, "GAME_FILE"):
        _mod.GAME_FILE = _test_game_file
    _orig_xlsx_const = getattr(_mod, "XLSX_FILE", None)
    if _orig_xlsx_const is None:
        continue
    # powergauge stores a str (os.path.join), the others a Path — preserve each.
    _mod.XLSX_FILE = _test_xlsx_path if isinstance(_orig_xlsx_const, Path) else str(_test_xlsx_path)

# Market-data cache constants: AETHER_CACHE_DIR (set at the top) covers modules imported
# AFTER this harness runs, but `unittest discover -s tests` imports test modules as
# top-level names, so this package can run only after earlier test modules already
# imported these with the real paths. Rebind them too (same object type as the original).
# (module, attribute, resolver, is_file): a file constant keeps its own file name.
for _mod_name, _attr, _resolve, _is_file in (
        ("aether.risk_utils", "OHLCV_DIR", "ohlcv_dir", False),
        ("rapidapi", "OHLCV_DIR", "ohlcv_dir", False),
        ("powergauge", "OHLCV_DIR", "ohlcv_dir", False),
        ("ai_portfolio_game", "SYMBOL_FULL_DIR", "ohlcv_dir", False),
        ("backtest_ratings", "OHLCV_DIR", "ohlcv_dir", False),
        ("backtest_ratings", "SYM_DIR", "symbol_dir", False),
        ("aether.circuit_breaker", "SPY_FILE", "ohlcv_dir", True),
        ("aether.circuit_breaker", "VXX_FILE", "ohlcv_dir", True),
        ("aether.scoring", "_OHLCV_ROOT", "ohlcv_dir", False),
        ("aether.decision_eval", "OHLCV_DIR", "ohlcv_dir", False),
        ("data_api", "_OHLCV_DIR", "ohlcv_dir", False),
        ("data_api", "_SYMBOL_DIR", "symbol_dir", False)):
    try:
        _mod = _importlib.import_module(_mod_name)
    except Exception:
        continue
    _orig = getattr(_mod, _attr, None)
    if _orig is None:
        continue
    _new = getattr(_paths, _resolve)()
    if _is_file:
        _new = str(Path(_new) / Path(_orig).name)
    setattr(_mod, _attr, Path(_new) if isinstance(_orig, Path) else _new)

# ---------------------------------------------------------------------------
# Learning-ledger guard — no test may write the decision log, the trade-DNA
# ledger, or the failure-DNA rules. Test fixtures (P1..P6, TSCO cost 100/stop 80)
# had leaked into these files, skewing the decision_eval scorecard and the
# retrospective analyzer. Unconditional, like the workbook guard.
# ---------------------------------------------------------------------------
import aether.decision_eval as _decision_eval
import aether.ledgers as _ledgers

_test_ledger_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
_ledger_tmp = Path(_test_ledger_dir.name)
_decision_eval.LOG = _ledger_tmp / "decision_log.jsonl"
_ledgers.TRADE_DNA_FILE = _ledger_tmp / "trade_history_dna.json"
_ledgers.FAILURE_RULES_FILE = _ledger_tmp / "failure_dna_rules.json"
_ledgers.RETRO_REPORT_FILE = _ledger_tmp / "retrospective_report.txt"

# ---------------------------------------------------------------------------
# Singleton-lock / trash / run-guard guard. run_watchdog() and
# autonomous_pipeline.main() overwrite their PID lock (tests mock "is the holder
# alive?" to False) and force-delete it at exit, and run_watchdog() purges the
# trash — so a suite run deleted the REAL Data/watchdog_run.lock, pipeline_run.lock
# and every >30-day file in Data/.trash, and could steal a live watchdog's or
# pipeline's lock. Unconditional: live tests never need the real locks either.
# Guarded by tests/test_log_hermetic.py::TestLocksTrashAndConfigRedirected.
# ---------------------------------------------------------------------------
import aether.config as _config
import aether.run_guard as _run_guard
import aether.trash as _trash

_test_lock_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
_lock_tmp = Path(_test_lock_dir.name)
_trash.TRASH_DIR = str(_lock_tmp / ".trash")
_run_guard._DATA_DIR = _lock_tmp
for _mod_name, _attrs in (("watchdog", ("WATCHDOG_LOCK_FILE", "SELF_HEAL_LOCK", "SELF_HEAL_PROMPT_FILE",
                                       "DATA_SENTINEL_FILE", "BACKUP_STATUS_FILE", "DATA_ALERT_MARKER")),
                          ("autonomous_pipeline", ("PIPELINE_LOCK_FILE",))):
    try:
        _mod = _importlib.import_module(_mod_name)
    except Exception:
        continue
    for _attr in _attrs:
        if hasattr(_mod, _attr):
            setattr(_mod, _attr, _lock_tmp / Path(getattr(_mod, _attr)).name)

if not _os.getenv("AETHER_LIVE_TESTS"):
    # -- Config side: never load the real config.json. Its E*TRADE TOTP secret opens
    # get_tokens' automated login path, so the suite behaved differently in the main
    # checkout than in a worktree (which has none). Rebuild CFG IN PLACE from a missing
    # file: modules that did `from aether.config import CFG` see the same object, and
    # nothing copies CFG values at import. Env-var overrides still apply, as in prod.
    _config._CFG_PATH = str(_lock_tmp / "config.json")  # deliberately absent
    _config.CFG.__init__()

    # -- State side: redirect prod auth-state / cache files to a throwaway temp dir --
    # A test that reaches get_tokens() finds no saved browser state (so the automated
    # headless path is skipped entirely — never even attempted) and any breaker
    # bookkeeping lands in temp, never polluting the production files. Tests that need
    # their own state file still patch these constants per-test (that override wins).
    _test_etrade_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
    _etrade._TOKEN_PATH         = str(Path(_test_etrade_dir.name) / "etrade_tokens.json")
    _etrade._BROWSER_STATE_PATH = str(Path(_test_etrade_dir.name) / "etrade_browser_state.json")
    _etrade._REAUTH_STATE_PATH  = str(Path(_test_etrade_dir.name) / "etrade_reauth_state.json")
    _etrade._REAUTH_LOCK_PATH   = str(Path(_test_etrade_dir.name) / "etrade_reauth.lock")

    # Scarcity-classification cache → temp, so a buy-path test that classifies a real
    # symbol never writes the production Data/scarcity_cache.json.
    _instruments._SCARCITY_CACHE_FILE = Path(_test_etrade_dir.name) / "scarcity_cache.json"

    # -- Wire side: block outbound sockets (except loopback) + Playwright launches --
    _ALLOWED_HOSTS = {"127.0.0.1", "localhost", "::1", "0.0.0.0"}

    def _host_of(address):
        # AF_INET -> (host, port); AF_UNIX/other -> str/path (local, allow).
        if isinstance(address, tuple):
            return address[0]
        return None

    def _blocked_net(host):
        raise RuntimeError(
            f"Blocked live network in tests (attempted connect to {host!r}). "
            "This would contact a production host — a real unmocked call once drove "
            "an Akamai bot-block toward an E*TRADE IP ban. Mock the network boundary, "
            "or set AETHER_LIVE_TESTS=1 to run the explicitly live-gated contract tests."
        )

    _orig_connect    = _socket.socket.connect
    _orig_connect_ex = _socket.socket.connect_ex

    def _guard_connect(self, address, *a, **k):
        host = _host_of(address)
        if host is None or host in _ALLOWED_HOSTS:
            return _orig_connect(self, address, *a, **k)
        _blocked_net(host)

    def _guard_connect_ex(self, address, *a, **k):
        host = _host_of(address)
        if host is None or host in _ALLOWED_HOSTS:
            return _orig_connect_ex(self, address, *a, **k)
        _blocked_net(host)

    _socket.socket.connect    = _guard_connect
    _socket.socket.connect_ex = _guard_connect_ex

    # Forbid launching a real browser under tests (E*TRADE / Chaikin auth path).
    if _pw_sync is not None:
        def _blocked_playwright(*_a, **_k):
            raise RuntimeError(
                "Blocked Playwright browser launch in tests (E*TRADE/Chaikin auth). "
                "Mock the browser path, or set AETHER_LIVE_TESTS=1."
            )

        _pw_sync.sync_playwright = _blocked_playwright
        
        # Apply the blocks directly to the module-level local references inside etrade and powergauge
        try:
            _etrade.sync_playwright = _blocked_playwright
        except (ImportError, AttributeError):
            pass
        try:
            _powergauge = _importlib.import_module("powergauge")
            _powergauge.sync_playwright = _blocked_playwright
        except (ImportError, AttributeError):
            pass


    # -- Process side: no test may kill a real process or launch the real game --
    # watchdog.run_watchdog() runs the live process supervisor (taskkill of duplicate
    # server.py / "orphaned" AETHER consoles) and a real `ai_portfolio_game.py --report`
    # child, which logs to the repo's Data/logs because a child never loads this harness.
    # A test that mocks subprocess itself still wins (its patch replaces these wrappers).

    # Kills, the real game, and anything that CHANGES Task Scheduler. Read-only
    # queries (`schtasks /query`, Get-ScheduledTask) stay allowed.
    # Matched on the command's shape, not a fixed string, so `schtasks.exe /Create`, a
    # full `C:\...\schtasks.exe` path and a quoted path are all caught.
    _FORBIDDEN_PROC_PATTERNS = (
        _re.compile(r"\btaskkill(\.exe)?\b"),
        _re.compile(r"\bstop-process\b"),
        _re.compile(r"ai_portfolio_game\.py"),
        _re.compile(r"\bschtasks(\.exe)?[\"']?\s+/(create|delete|change|run|end)\b"),
        _re.compile(r"\b(register|unregister|set|start|stop|enable|disable)-scheduledtask\b"),
    )

    def _argv_text(args):
        if isinstance(args, (list, tuple)):
            return " ".join(str(a) for a in args)
        return str(args)

    def _guard_proc(args):
        text = _argv_text(args).lower()
        hit = next((p.pattern for p in _FORBIDDEN_PROC_PATTERNS if p.search(text)), None)
        if hit:
            raise RuntimeError(
                f"Blocked real process side effect in tests ({hit!r} in {_argv_text(args)[:120]!r}). "
                "Mock subprocess at the boundary, or set AETHER_LIVE_TESTS=1."
            )

    _orig_subprocess_run = _subprocess.run
    _orig_subprocess_popen = _subprocess.Popen

    def _guarded_run(args, *a, **k):
        _guard_proc(args)
        return _orig_subprocess_run(args, *a, **k)

    class _GuardedPopen(_orig_subprocess_popen):
        def __init__(self, args, *a, **k):
            _guard_proc(args)
            super().__init__(args, *a, **k)

    _subprocess.run = _guarded_run
    _subprocess.Popen = _GuardedPopen
